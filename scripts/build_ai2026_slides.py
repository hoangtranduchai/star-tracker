#!/usr/bin/env python3
"""Build AI2026 slides from HD1024 thr=2 GPU benchmark ONLY."""

from __future__ import annotations

import json
import sys
from pathlib import Path

import matplotlib

matplotlib.use("Agg")
import matplotlib.pyplot as plt
from matplotlib.patches import Ellipse, FancyBboxPatch, Rectangle
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.util import Inches, Pt

ROOT = Path(__file__).resolve().parent.parent
OUT = ROOT / "AI2026_SLIDE_THUYET_TRINH.pptx"
FIG = ROOT / "data" / "outputs" / "ai2026_figures"
BENCH = ROOT / "data" / "outputs" / "centroid_benchmark_hd1024_thr2.json"
GLARE = ROOT / "data" / "outputs" / "centroid_benchmark_hd1024_glare_sun_strong.json"
MOUNT = ROOT / "data" / "outputs" / "mount_st_helens_centroid_compare.json"
MOUNT_LOST = ROOT / "data" / "outputs" / "mount_st_helens_lost_compare.json"
MOUNT_OVERLAY = ROOT / "data" / "outputs" / "mount_st_helens_cog_vs_cnn.png"
ZHAO_FLOW = ROOT / "data" / "outputs" / "zhao_paper_figures" / "project_00.png"
ZHAO_STRAY = ROOT / "data" / "outputs" / "zhao_paper_figures" / "project_03.png"
PREVIEW = ROOT / "data" / "sim_dataset" / "hd1024" / "0001_preview.png"

NAVY = RGBColor(0x1F, 0x4E, 0x79)
DARK = RGBColor(0x22, 0x22, 0x22)
MUTED = RGBColor(0x55, 0x55, 0x55)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)


def _run(p, text, size=18, bold=False, color=DARK):
    r = p.add_run()
    r.text = text
    r.font.size = Pt(size)
    r.font.bold = bold
    r.font.color.rgb = color
    r.font.name = "Calibri"


def title_bar(slide, title, subtitle=None):
    box = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), Inches(13.333), Inches(0.95))
    box.fill.solid()
    box.fill.fore_color.rgb = NAVY
    box.line.fill.background()
    tf = box.text_frame
    tf.clear()
    _run(tf.paragraphs[0], title, size=26, bold=True, color=WHITE)
    if subtitle:
        _run(tf.add_paragraph(), subtitle, size=12, color=RGBColor(0xD6, 0xE4, 0xF0))


def bullets(slide, left, top, width, height, items, size=16):
    box = slide.shapes.add_textbox(Inches(left), Inches(top), Inches(width), Inches(height))
    tf = box.text_frame
    tf.word_wrap = True
    for i, item in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.space_after = Pt(8)
        _run(p, "• " + item, size=size, color=DARK)


def footer(slide, page, total):
    box = slide.shapes.add_textbox(Inches(0.4), Inches(7.1), Inches(12.5), Inches(0.35))
    _run(
        box.text_frame.paragraphs[0],
        f"Trường Đại học Bách khoa, Đại học Đà Nẵng — Khoa Công nghệ thông tin    |    {page}/{total}",
        size=10,
        color=MUTED,
    )


def blank(prs):
    return prs.slides.add_slide(prs.slide_layouts[6])


def fig_pipeline():
    FIG.mkdir(parents=True, exist_ok=True)
    fig, ax = plt.subplots(figsize=(12.5, 3.4), dpi=160)
    ax.set_xlim(0, 12.5)
    ax.set_ylim(0, 3.4)
    ax.axis("off")
    boxes = [
        (0.3, 1.0, 2.0, 1.4, "Ảnh sao\n1024×683"),
        (2.6, 1.0, 2.0, 1.4, "Centroid\nCoG | CNN*"),
        (4.9, 1.0, 2.2, 1.4, "Pyramid +\nK-vector"),
        (7.4, 1.0, 2.0, 1.4, "Wahba\nSVD/QUEST"),
        (9.7, 1.0, 2.4, 1.4, "Quaternion\n16 byte"),
    ]
    for x, y, w, h, label in boxes:
        ax.add_patch(
            FancyBboxPatch(
                (x, y), w, h, boxstyle="round,pad=0.04,rounding_size=0.12",
                facecolor="#e8f1fb", edgecolor="#1f4e79", lw=1.6,
            )
        )
        ax.text(x + w / 2, y + h / 2, label, ha="center", va="center", fontsize=11)
    for x0, x1 in [(2.3, 2.6), (4.6, 4.9), (7.1, 7.4), (9.4, 9.7)]:
        ax.annotate("", xy=(x1, 1.7), xytext=(x0, 1.7),
                    arrowprops=dict(arrowstyle="->", color="#1f4e79", lw=1.8))
    ax.text(
        6.25, 0.3,
        "* CNN MobileUNet công khai · d_th=2 (public) · không huấn luyện lại",
        ha="center", fontsize=8.5, color="#444",
    )
    ax.set_title("Đường ống star tracker (demo laptop)", fontsize=13, color="#1f4e79")
    path = FIG / "pipeline.png"
    fig.tight_layout()
    fig.savefig(path, bbox_inches="tight", facecolor="white")
    plt.close(fig)
    return path


def fig_camera():
    fig, ax = plt.subplots(figsize=(10, 5.0), dpi=160)
    ax.set_xlim(0, 10)
    ax.set_ylim(0, 5.0)
    ax.axis("off")
    ax.add_patch(Rectangle((0.6, 1.5), 2.4, 1.5, fill=False, lw=2, edgecolor="#1f4e79"))
    ax.text(1.8, 2.25, "APS-C demo\n6000×4000\np=3,72 µm", ha="center", va="center", fontsize=9)
    ax.add_patch(Ellipse((4.0, 2.25), 0.55, 1.4, fill=False, lw=2, edgecolor="#c45c26"))
    ax.text(4.0, 0.7, "f = 50 mm", ha="center", fontsize=10, color="#c45c26")
    ax.plot([4.3, 8.8], [3.0, 4.4], color="#1f4e79", lw=1.2)
    ax.plot([4.3, 8.8], [1.5, 0.5], color="#1f4e79", lw=1.2)
    ax.plot([4.3, 8.8], [2.25, 2.25], color="#888", lw=1.0, ls="--")
    ax.text(7.0, 4.6, "FOV ≈ 25,16° × 16,93°", fontsize=10, color="#1f4e79")
    ax.add_patch(
        FancyBboxPatch((0.5, 3.4), 3.2, 1.3, boxstyle="round,pad=0.05",
                       facecolor="#fff4e8", edgecolor="#c45c26", lw=1.4)
    )
    ax.text(2.1, 4.05, "Benchmark hd1024\n1024×683\n(cùng lớp FOV demo)", ha="center", va="center", fontsize=9)
    ax.add_patch(
        FancyBboxPatch((5.4, 1.4), 4.1, 2.2, boxstyle="round,pad=0.05",
                       facecolor="#f3f7fb", edgecolor="#1f4e79", lw=1.2)
    )
    ax.text(
        7.45, 2.5,
        "Demo đầy đủ: 15,35″/pixel\nTâm: u₀=2999,5  v₀=1999,5\n\n"
        "hd1024 = tỷ lệ cạnh dài ~1024\nkhông phải cảm biến bay mới",
        ha="center", va="center", fontsize=9,
    )
    ax.set_title("Hình học camera demo / bộ đo 1024", fontsize=12, color="#1f4e79")
    path = FIG / "camera_geometry.png"
    fig.tight_layout()
    fig.savefig(path, bbox_inches="tight", facecolor="white")
    plt.close(fig)
    return path


def fig_starfield():
    path = FIG / "starfield_hd1024.png"
    src = PREVIEW if PREVIEW.is_file() else None
    if src is None:
        cands = sorted((ROOT / "data" / "sim_dataset" / "hd1024").glob("*_preview.png"))
        src = cands[1] if len(cands) > 1 else (cands[0] if cands else None)
    if src is None:
        return None
    img = plt.imread(str(src))
    fig, ax = plt.subplots(figsize=(10, 5.0), dpi=150)
    ax.imshow(img, cmap="gray" if getattr(img, "ndim", 2) == 2 else None)
    ax.set_title("Ảnh sao benchmark 1024×683 (preview từ data/sim_dataset/hd1024)")
    ax.axis("off")
    fig.tight_layout()
    fig.savefig(path, bbox_inches="tight", facecolor="white")
    plt.close(fig)
    return path


def fig_comparison(s: dict):
    cl, nn = s["classical"], s["cnn"]
    thr = s.get("distance_map_threshold", 2.0)
    rows = [
        ["Phương án", "RMSE tb", "RMSE tv", "Khớp/GT", "ms tv", "s/khung tb"],
        [
            "CoG cổ điển",
            f"{cl['rmse_px_mean']:.4f}",
            f"{cl['rmse_px_median']:.4f}",
            f"{cl['matched_stars_total']}/{cl['gt_stars_total']}",
            f"{cl['time_ms_median']:.2f}",
            f"{cl['time_s_per_frame_mean']:.3f}",
        ],
        [
            f"CNN MobileUNet (d_th={thr:g})",
            f"{nn['rmse_px_mean']:.4f}",
            f"{nn['rmse_px_median']:.4f}",
            f"{nn['matched_stars_total']}/{nn['gt_stars_total']}",
            f"{nn['time_ms_median']:.2f}",
            f"{nn['time_s_per_frame_mean']:.3f}",
        ],
    ]
    fig, ax = plt.subplots(figsize=(12, 2.7), dpi=160)
    ax.axis("off")
    table = ax.table(cellText=rows, loc="center", cellLoc="center")
    table.auto_set_font_size(False)
    table.set_fontsize(10)
    table.scale(1.15, 1.7)
    for j in range(6):
        table[0, j].set_facecolor("#1f4e79")
        table[0, j].set_text_props(color="white", weight="bold")
    ax.set_title(
        "So sánh centroid — 500 khung 1024×683 · d_th=2 · RTX 5070 (đo lại; lần ngưỡng sai đã loại)",
        fontsize=11, pad=10, color="#1f4e79",
    )
    path = FIG / "comparison_table.png"
    fig.tight_layout()
    fig.savefig(path, bbox_inches="tight", facecolor="white")
    plt.close(fig)
    return path


def build():
    s = json.loads(BENCH.read_text(encoding="utf-8")).get("summary")
    if not s or s.get("classical", {}).get("n_frames") != 500:
        raise SystemExit("thr2 JSON incomplete")
    cl, nn, cam = s["classical"], s["cnn"], s["camera"]
    thr = s.get("distance_map_threshold", 2.0)
    m = json.loads(MOUNT.read_text(encoding="utf-8")) if MOUNT.is_file() else None
    g = json.loads(GLARE.read_text(encoding="utf-8")).get("summary") if GLARE.is_file() else None

    fig_pipeline()
    fig_camera()
    fig_starfield()
    fig_comparison(s)

    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    total = 13

    sl = blank(prs)
    bg = sl.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), Inches(13.333), Inches(7.5))
    bg.fill.solid()
    bg.fill.fore_color.rgb = NAVY
    bg.line.fill.background()
    tb = sl.shapes.add_textbox(Inches(0.8), Inches(1.8), Inches(11.5), Inches(4.5))
    tf = tb.text_frame
    tf.word_wrap = True
    _run(tf.paragraphs[0], "STAR TRACKER — DEMO LAPTOP", size=34, bold=True, color=WHITE)
    p = tf.add_paragraph()
    _run(p, "Centroid → Pyramid/K-vector → Wahba → Quaternion 16 byte", size=18, color=RGBColor(0xD6, 0xE4, 0xF0))
    p = tf.add_paragraph()
    p.space_before = Pt(20)
    _run(p, "Cuộc thi Sáng tạo trẻ Quốc gia — Trí tuệ nhân tạo 2026 (Bảng C)", size=15, color=WHITE)
    p = tf.add_paragraph()
    p.space_before = Pt(16)
    _run(p, "Trường Đại học Bách khoa, Đại học Đà Nẵng\nKhoa Công nghệ thông tin", size=18, bold=True, color=WHITE)
    p = tf.add_paragraph()
    p.space_before = Pt(12)
    _run(
        p,
        "Đội: Hoàng Trần Đức Hải · Nguyễn Tiến · Mai Tạ Trúc Vy\n"
        "GVHD: TS. Nguyễn Năng Hùng Vân",
        size=12,
        color=RGBColor(0xB8, 0xC9, 0xDC),
    )

    sl = blank(prs)
    title_bar(sl, "Nội dung thuyết trình")
    bullets(sl, 0.7, 1.3, 11.5, 5.5, [
        "Bài toán lost-in-space và quaternion 16 byte",
        "Pipeline cổ điển + CNN centroid tùy chọn (d_th=2)",
        "Kết quả đo hd1024: CoG thắng RMSE trên bộ sạch",
        "Điểm sáng tạo: ảnh núi thật LOST — CoG sao giả / CNN lọc",
        "Trích dẫn Zhao et al. (arXiv:2404.19108) — tách khỏi số đo nhóm",
        "Giới hạn & FAQ",
    ], size=18)
    footer(sl, 2, total)

    sl = blank(prs)
    title_bar(sl, "1. Bài toán lost-in-space")
    bullets(sl, 0.7, 1.3, 11.5, 5.5, [
        "Không biết hướng nhìn ban đầu trên bầu trời.",
        "Đầu ra: quaternion J2000 — 4×float32 = 16 byte.",
        "Giải bằng hình học góc giữa các sao.",
        "Demo mặc định không ML; CNN chỉ tùy chọn cho centroid.",
    ], size=18)
    footer(sl, 3, total)

    sl = blank(prs)
    title_bar(sl, "2. Đường ống thuật toán")
    img = FIG / "pipeline.png"
    if img.is_file():
        sl.shapes.add_picture(str(img), Inches(0.5), Inches(1.25), width=Inches(12.3))
    bullets(sl, 0.7, 5.0, 12, 1.8, [
        "Pyramid (Mortari) + K-vector · Yale BSC5 V ≤ 6.",
        "Wahba SVD/QUEST → quaternion.",
        f"CNN: MobileUNet công khai, d_th={thr:g} (public), không huấn luyện lại.",
    ], size=14)
    footer(sl, 4, total)

    sl = blank(prs)
    title_bar(sl, "3. Hình học camera / bộ đo 1024")
    img = FIG / "camera_geometry.png"
    if img.is_file():
        sl.shapes.add_picture(str(img), Inches(1.0), Inches(1.15), width=Inches(11.2))
    footer(sl, 5, total)

    sl = blank(prs)
    title_bar(sl, "4. Ảnh sao benchmark hd1024")
    img = FIG / "starfield_hd1024.png"
    if img.is_file():
        sl.shapes.add_picture(str(img), Inches(1.5), Inches(1.2), width=Inches(10.3))
    elif PREVIEW.is_file():
        sl.shapes.add_picture(str(PREVIEW), Inches(2.5), Inches(1.3), height=Inches(5.2))
    footer(sl, 6, total)

    sl = blank(prs)
    title_bar(sl, "5. Kết quả đo (d_th=2)", "centroid_benchmark_hd1024_thr2.json — lần ngưỡng sai đã loại")
    img = FIG / "comparison_table.png"
    if img.is_file():
        sl.shapes.add_picture(str(img), Inches(0.5), Inches(1.15), width=Inches(12.2))
    bullets(sl, 0.7, 4.4, 12, 2.4, [
        f"500 khung {cam['width']}×{cam['height']}, tái dùng ảnh có sẵn, GPU {s.get('device')}.",
        f"CoG: RMSE tb {cl['rmse_px_mean']:.4f} px · {cl['time_ms_median']:.2f} ms tv "
        f"(~{cl['time_s_per_frame_mean']:.3f} s/khung).",
        f"CNN: RMSE tb {nn['rmse_px_mean']:.4f} px · {nn['time_ms_median']:.2f} ms tv "
        f"(~{nn['time_s_per_frame_mean']:.3f} s/khung); khớp {nn['matched_stars_total']}/{nn['gt_stars_total']}.",
        "Trên bộ sạch này CoG vừa nhanh hơn vừa chính xác hơn CNN công khai.",
        (
            f"Lóa mạnh {g['classical']['n_frames']} khung: CoG RMSE {g['classical']['rmse_px_mean']:.4f} px "
            f"vs CNN {g['cnn']['rmse_px_mean']:.4f} px — CNN không thắng RMSE."
            if g
            else "Lóa mạnh: xem centroid_benchmark_hd1024_glare_sun_strong.json."
        ),
        "Lần d_th≈0,707 đã loại. Cổng CNN «cấp 6 trở lên sáng»: có trên web; 1 khung sim thử không thắng CoG.",
    ], size=12)
    footer(sl, 7, total)

    # --- creativity: mountain photo ---
    sl = blank(prs)
    title_bar(
        sl,
        "6. Điểm sáng tạo — ảnh núi thật LOST",
        "mount_st_helens_1.png · 4,2 mm · 4,1 µm · không phải ảnh nhóm chụp",
    )
    if MOUNT_OVERLAY.is_file():
        sl.shapes.add_picture(str(MOUNT_OVERLAY), Inches(0.4), Inches(1.15), height=Inches(5.5))
    if m is not None:
        lost = None
        if MOUNT_LOST.is_file():
            lost = json.loads(MOUNT_LOST.read_text(encoding="utf-8"))
        modes = (lost or {}).get("modes") or {}
        cnn_a = modes.get("cnn") or {}
        cog_a = modes.get("classical") or {}
        att_lines = []
        if cnn_a.get("ra_deg") is not None:
            att_lines.append(
                f"CNN web: RA {cnn_a['ra_deg']:.2f}° Dec {cnn_a['dec_deg']:+.2f}° "
                f"(Δ {cnn_a['sep_deg_from_LOST']:.3f}° vs LOST 310.446/+36.023) · "
                f"ID {cnn_a.get('n_matched')}/{cnn_a.get('n_detected')} · "
                f"~{float(cnn_a.get('pipeline_ms') or 0)/1000.0:.1f} s"
            )
        if cog_a.get("ra_deg") is not None:
            att_lines.append(
                f"CoG web: RA {cog_a['ra_deg']:.2f}° Dec {cog_a['dec_deg']:+.2f}° "
                f"(Δ {cog_a['sep_deg_from_LOST']:.3f}° vs LOST) · "
                f"ID {cog_a.get('n_matched')}/{cog_a.get('n_detected')} · "
                f"~{float(cog_a.get('pipeline_ms') or 0)/1000.0:.1f} s"
            )
        bullets(sl, 7.6, 1.3, 5.3, 5.5, [
            f"CoG: {m['classical']['n_total']} điểm · núi/cây {m['classical']['n_on_terrain']} · trời {m['classical']['n_on_sky']}",
            f"CNN: {m['cnn']['n_total']} điểm · núi/cây {m['cnn']['n_on_terrain']} · trời {m['cnn']['n_on_sky']}",
            "Cam = CoG đánh sao giả trên núi/cây; xanh = CNN gần như chỉ trên trời.",
            *att_lines,
            "Web khóa: CNN ~5 s · 63/79 · RA 310,49° Dec +36,01°; CoG ~16 s · 64/106 · RA 310,52° Dec +35,95° (cả hai ≲0,1° vs LOST).",
            "CNN nhanh hơn trên ca núi vì không nhai sao giả địa hình — không claim luôn tốt hơn.",
            "Sim sạch: CoG ~0,446″ vs CNN ~1,65″ — CoG vẫn chính xác hơn.",
            "Video minh chứng + git remote: chưa kiểm được.",
        ], size=11)
    footer(sl, 8, total)

    # --- Zhao citations ---
    sl = blank(prs)
    title_bar(sl, "7. Trích dẫn Zhao et al.", "arXiv:2404.19108 — số liệu TÁC GIẢ, không phải đo nhóm")
    if ZHAO_FLOW.is_file():
        sl.shapes.add_picture(str(ZHAO_FLOW), Inches(0.4), Inches(1.2), width=Inches(6.2))
    if ZHAO_STRAY.is_file():
        sl.shapes.add_picture(str(ZHAO_STRAY), Inches(6.9), Inches(1.2), width=Inches(5.8))
    bullets(sl, 0.5, 5.0, 12.3, 2.0, [
        "Nguồn: Zhao et al., arXiv:2404.19108 — flowchart (trái) · night+stray (phải).",
        "Table V: MobileUNet ≈ 0,020 s (~20 ms) trên RTX 2060M @ 640×480; Table VI: Coral TPU ≈ 266 ms.",
        "Table III: RMSE MobileUNet ≈ 0,1695 px vs CoG ≈ 0,6966 px (ảnh tổng hợp của tác giả).",
        f"Đo nhóm (RTX 5070, 1024×683, d_th=2): CoG {cl['rmse_px_mean']:.4f} px / ~34 ms; CNN {nn['rmse_px_mean']:.4f} px / ~60 ms.",
        "Caption hình: Nguồn: Zhao et al., arXiv:2404.19108 — không phải đo của nhóm.",
    ], size=11)
    footer(sl, 9, total)

    sl = blank(prs)
    title_bar(sl, "8. Kiến trúc & chạy thử")
    bullets(sl, 0.7, 1.3, 11.5, 5.5, [
        "camera → centroid (CoG|CNN) → Pyramid/K-vector → Wahba → quaternion 16 byte",
        "CLI: python demo_star_tracker.py --mode fast --seed 42 --lookup --no-show",
        "Web: python app.py → preset «Ảnh núi thật» / «Preset lóa mạnh»",
        "JSON: centroid_benchmark_hd1024_thr2.json · mount_st_helens_centroid_compare.json · mount_st_helens_lost_compare.json",
        "Thiếu: hai file video minh chứng; URL git remote — chưa kiểm được.",
    ], size=16)
    footer(sl, 10, total)

    sl = blank(prs)
    title_bar(sl, "9. Giới hạn")
    bullets(sl, 0.7, 1.3, 11.5, 5.5, [
        "Demo laptop — không phải thiết bị bay.",
        "Trên hd1024 sạch: CNN công khai lệch miền; RMSE cao hơn CoG.",
        "Mặt nạ núi/cây là heuristic — kèm overlay để kiểm chứng mắt.",
        "Attitude end-to-end (Δθ) trên bộ sim lớn: chưa kiểm được.",
    ], size=17)
    footer(sl, 11, total)

    sl = blank(prs)
    title_bar(sl, "10. FAQ")
    mount_faq = "chưa kiểm được"
    if m is not None:
        mount_faq = (
            f"có — CoG núi/cây {m['classical']['n_on_terrain']} vs CNN {m['cnn']['n_on_terrain']}; "
            f"web CNN 63/79 ~5 s RA 310,49° Dec +36,01°; CoG 64/106 ~16 s RA 310,52° Dec +35,95° "
            f"(cả hai ≲0,1° vs LOST 310,446/+36,023)."
        )
    bullets(sl, 0.7, 1.3, 11.5, 5.5, [
        "Có AI không? — Demo mặc định không ML; CNN centroid tùy chọn công khai.",
        f"CNN có tốt hơn? — Trên 500 khung hd1024: chưa (RMSE). Trên ảnh núi thật: {mount_faq}",
        "Đầu ra? — Quaternion 16 byte (J2000).",
        "Catalog? — Yale BSC5, V ≤ 6.",
    ], size=16)
    footer(sl, 12, total)

    sl = blank(prs)
    title_bar(sl, "Kết · Cảm ơn quý thầy cô")
    bullets(sl, 0.7, 1.4, 11.5, 5, [
        "Hệ lai + so sánh công khai + ca khó ảnh núi thật.",
        f"hd1024 (d_th=2): CoG RMSE tốt hơn CNN công khai trên bộ sạch.",
        "Ảnh núi LOST: CNN lọc sao giả trên núi/cây tốt hơn CoG (đo thật).",
        "Hồ sơ: AI2026_Ho-so-du-an_StarTracker.docx / .md",
        "GVHD: TS. Nguyễn Năng Hùng Vân",
        "Trường Đại học Bách khoa, Đại học Đà Nẵng — Khoa Công nghệ thông tin",
    ], size=16)
    footer(sl, 13, total)

    prs.save(str(OUT))
    print(f"wrote {OUT}")


if __name__ == "__main__":
    build()
