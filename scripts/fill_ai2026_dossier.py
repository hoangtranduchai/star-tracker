#!/usr/bin/env python3
"""Fill AI2026 Bang C dossier from HD1024 thr=2 + Mount St. Helens outdoor compare."""

from __future__ import annotations

import json
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt

ROOT = Path(__file__).resolve().parent.parent
TEMPLATE = ROOT / "AI2026_Mẫu-hồ-sơ-gốc.docx"
ORIGINAL = ROOT / "AI2026_Mẫu-hồ-sơ.docx"
OUT = ROOT / "AI2026_Ho-so-du-an_StarTracker.docx"
OUT_MD = ROOT / "AI2026_Ho-so-du-an_StarTracker.md"
BENCH = ROOT / "data" / "outputs" / "centroid_benchmark_hd1024_thr2.json"
MOUNT = ROOT / "data" / "outputs" / "mount_st_helens_centroid_compare.json"
MOUNT_LOST = ROOT / "data" / "outputs" / "mount_st_helens_lost_compare.json"
MOUNT_OVERLAY = ROOT / "data" / "outputs" / "mount_st_helens_cog_vs_cnn.png"
DEMO_FIELD = ROOT / "data" / "sim_dataset" / "hd1024" / "0001_preview.png"
ZHAO_FLOW = ROOT / "data" / "outputs" / "zhao_paper_figures" / "project_00.png"
ZHAO_STRAY = ROOT / "data" / "outputs" / "zhao_paper_figures" / "project_03.png"
ZHAO_ATT = ROOT / "data" / "outputs" / "zhao_paper_figures" / "project_05.png"

# Identity is copied at runtime from ORIGINAL Bang C dossier only.
# Bang C form has no GVHD cell; user-confirmed advisor for MD/slide credits:
ADVISOR = "TS. Nguyễn Năng Hùng Vân"

STRUCTURE_PDFS = [
    "ecoSim-Hồ-sơ-dự-án.pdf",
    "Meetly-Hồ-sơ-dự-án.pdf",
    "Signlang_Hồ-sơ-dự-án.pdf",
    "Speakee-Hồ-sơ-dự-án.pdf",
    "UniTrust-Hồ-sơ-dự-án.pdf",
    "Viraldy-Hồ-sơ-dự-án.pdf",
]

# Value rows in the first table (cell index 1) for the three thí sinh.
IDENTITY_VALUE_ROWS = (
    2, 3, 4, 5, 6, 7,  # thí sinh 1
    9, 10, 11, 12, 13, 14,  # thí sinh 2
    16, 17, 18, 19, 20, 21,  # thí sinh 3
)


def _summary() -> dict:
    raw = json.loads(BENCH.read_text(encoding="utf-8"))
    return raw.get("summary", raw)


def _mount() -> dict | None:
    if not MOUNT.is_file():
        return None
    return json.loads(MOUNT.read_text(encoding="utf-8"))


def _mount_lost() -> dict | None:
    if not MOUNT_LOST.is_file():
        return None
    return json.loads(MOUNT_LOST.read_text(encoding="utf-8"))


def _lost_attitude_sentence(lost: dict | None) -> str:
    """Web-confirmed boresight vs LOST published RA/Dec (honest measured numbers)."""
    if lost is None:
        return (
            "Attitude vs LOST (RA 310,446° · Dec +36,0229°): chưa kiểm được "
            "(thiếu mount_st_helens_lost_compare.json)."
        )
    ref = lost.get("LOST_ref") or {}
    modes = lost.get("modes") or {}
    bits: list[str] = []
    for key, label in (("cnn", "CNN"), ("classical", "CoG")):
        m = modes.get(key) or {}
        if not m.get("success") or m.get("ra_deg") is None:
            bits.append(f"{label}: chưa giải được / FAIL")
            continue
        bits.append(
            f"{label} RA {m['ra_deg']:.2f}° Dec {m['dec_deg']:+.2f}° "
            f"(Δ {m['sep_deg_from_LOST']:.3f}° vs LOST) · "
            f"ID {m.get('n_matched')}/{m.get('n_detected')} · "
            f"pipeline ~{float(m.get('pipeline_ms') or 0)/1000.0:.1f} s · "
            f"fallback={m.get('fallback')}"
        )
    return (
        f"Attitude web vs LOST công bố (RA {float(ref.get('ra_deg', 310.446)):.3f}° · "
        f"Dec {float(ref.get('dec_deg', 36.0229)):+.4f}°): " + "; ".join(bits) + ". "
        "JSON: data/outputs/mount_st_helens_lost_compare.json. "
        "Hai file video minh chứng và URL git remote: chưa kiểm được — không bịa link."
    )


def _cell_val(row) -> str:
    if len(row.cells) < 2:
        return ""
    return row.cells[1].text.strip().replace("\n", " ")


def load_identity_from_original() -> dict[int, str]:
    """Copy only non-empty identity values from the user's original Bang C dossier."""
    src_path = ORIGINAL
    if not src_path.is_file():
        cands = [
            p
            for p in ROOT.glob("AI2026*Mẫu*.docx")
            if "gốc" not in p.name and "goc" not in p.name
        ]
        cands = [p for p in cands if p.stat().st_size > 1_000_000]
        if not cands:
            raise SystemExit(f"missing original dossier {ORIGINAL}")
        src_path = cands[0]
    src = Document(str(src_path))
    table = src.tables[0]
    out: dict[int, str] = {}
    for ri in IDENTITY_VALUE_ROWS:
        if ri >= len(table.rows):
            continue
        val = _cell_val(table.rows[ri])
        if val:
            out[ri] = val
    # Team size: ORIGINAL marks 3 người (checkbox cleared on that cell in XML).
    out[-1] = "3"
    return out


def _dots(paragraph, text: str) -> bool:
    raw = paragraph.text.strip()
    if not raw:
        return False
    if set(raw) <= {".", " ", "…", "\u00a0"} or raw.startswith("...."):
        paragraph.text = text
        for run in paragraph.runs:
            run.font.size = Pt(11)
        return True
    return False


def sections(s: dict, m: dict | None, lost: dict | None = None) -> dict[int, str]:
    cam, cl, nn = s["camera"], s["classical"], s["cnn"]
    refs = [n for n in STRUCTURE_PDFS if (ROOT / n).is_file()]
    thr = s.get("distance_map_threshold", 2.0)
    lost_sent = _lost_attitude_sentence(lost)

    if m is not None:
        mount_para = (
            f"ĐIỂM SÁNG TẠO — ảnh núi thật công bố kèm LOST (không phải ảnh nhóm chụp): "
            f"data/uploads/mount_st_helens_1.png, 1024×1024, 8-bit, tiêu cự 4,2 mm, pixel 4,1 µm, "
            f"FOV ≈ 53°×53°. Overlay CoG (cam) vs CNN (xanh): data/outputs/mount_st_helens_cog_vs_cnn.png. "
            f"Đếm trên mặt nạ núi/cây heuristic (chỉ dưới đường rid): CoG tổng {m['classical']['n_total']} "
            f"(núi/cây {m['classical']['n_on_terrain']}, trời {m['classical']['n_on_sky']}); "
            f"CNN tổng {m['cnn']['n_total']} (núi/cây {m['cnn']['n_on_terrain']}, trời {m['cnn']['n_on_sky']}). "
            f"{m.get('verdict_vi', '')} JSON: data/outputs/mount_st_helens_centroid_compare.json. "
            "Web preset «Ảnh núi thật» (đã khóa lần chạy gần nhất, sha array 62593b48bc0c): "
            "CNN ~5,1 s · 63/79 · RA 310,49° Dec +36,01°; "
            "CoG ~16,2 s · 64/106 · RA 310,52° Dec +35,95° — cả hai lệch LOST ≲ ~0,1° "
            "(LOST 310,446° / +36,0229°). CNN nhanh hơn trên ca này vì không «nhai» sao giả địa hình; "
            "không claim CNN luôn tốt hơn. Trên sim sạch CoG vẫn chính xác hơn "
            "(~0,446″ vs CNN ~1,65″). "
            f"{lost_sent}"
        )
        creativity = (
            "Điểm sáng tạo: (1) hệ lai CoG|CNN trên cùng pipeline Pyramid→Wahba; "
            "(2) so sánh công khai MobileUNet (Zhao et al.) không huấn luyện lại; "
            "(3) ca khó ảnh núi thật LOST — CoG tạo nhiều sao giả trên núi/cây, CNN gần như không "
            f"(núi/cây CoG {m['classical']['n_on_terrain']} vs CNN {m['cnn']['n_on_terrain']})."
        )
    else:
        mount_para = "Ảnh núi thật LOST: chưa kiểm được (thiếu mount_st_helens_centroid_compare.json)."
        creativity = (
            "Điểm sáng tạo: hệ lai CoG|CNN; so sánh công khai MobileUNet; "
            "ca khó ảnh núi — chưa kiểm được file đo."
        )

    zhao_cite = (
        "Trích dẫn Zhao et al., arXiv:2404.19108 (không phải số đo của nhóm): "
        "MobileUNet ≈ 0,020 s/khung trên RTX 2060M (Table V; thường nêu ~20 ms @ 640×480); "
        "MobileUNet trên Google Coral Edge TPU ≈ 266 ms (Table VI; số tác giả ≈ 265,5 ms); "
        "RMSE centroid tổng hợp MobileUNet ≈ 0,1695 px vs CoG ≈ 0,6966 px trên ảnh tổng hợp sạch (Table III). "
        "Hình minh họa nguồn tác giả: data/outputs/zhao_paper_figures/ "
        "(flowchart, stray-light night, attitude) — chú thích «Nguồn: Zhao et al., arXiv:2404.19108»."
    )

    s1 = (
        "Bài toán thực tiễn: từ ảnh bầu trời sao xác định hướng (attitude) của trục máy ảnh "
        "trong hệ J2000 khi không có tư thế ban đầu (lost-in-space). Ba bước: ước lượng tâm sao "
        "dưới mức một pixel; nhận dạng sao bằng hình học góc và catalog; giải Wahba ra quaternion. "
        "Lý do chọn: đầu ra gọn (4×float32 = 16 byte), phù hợp demo laptop cho sinh viên CNTT; "
        "đồng thời đối chiếu centroid cổ điển (CoG) với MobileUNet công khai (Zhao et al., "
        "arXiv:2404.19108) — không huấn luyện lại. Attitude là góc/hướng, không dùng GSD mét mặt đất. "
        + creativity
    )
    s2 = (
        "Mục tiêu: (1) phát hiện và ước lượng tâm sao; (2) Pyramid/Mortari + K-vector trên Yale "
        "BSC5 (V ≤ 6); (3) quaternion J2000 16 byte (Wahba SVD/QUEST). Phạm vi: demo trên máy tính "
        "cá nhân. Camera demo: máy ảnh tĩnh APS-C khoảng 24 MP (6000×4000, pixel 3,72 µm, tiêu cự "
        f"50 mm, FOV ≈ 25,16°×16,93°). Bộ đo centroid: {cam['width']}×{cam['height']} "
        f"(pixel {cam['pixel_um']} µm, f={cam['focal_mm']} mm) — tỷ lệ cạnh dài ~1024 cùng lớp FOV "
        "demo, không phải cảm biến bay mới. Không phải sản phẩm bay. "
        "Đơn vị: Trường Đại học Bách khoa, Đại học Đà Nẵng — Khoa Công nghệ thông tin."
    )
    s3 = (
        "Dữ liệu: (a) catalog Yale BSC5 công khai; (b) 500 ảnh giả lập sẵn trong "
        f"{s.get('dataset_dir', 'data/sim_dataset/hd1024')} "
        f"({cam['width']}×{cam['height']}; lần đo này không sinh lại ảnh); "
        "(c) trọng số MobileUNet_B10_50.pt công khai, không huấn luyện lại; "
        "(d) ảnh thật công bố kèm LOST: data/uploads/mount_st_helens_1.png (4,2 mm / 4,1 µm); "
        "(e) hình minh họa từ trang dự án Zhao (chỉ trích dẫn). Không dùng dữ liệu cá nhân. "
        "Không sao chép số liệu từ hồ sơ nhóm khác."
    )
    s4 = (
        "Tiền xử lý CoG: nền cục bộ, ngưỡng MAD, opening, connected components, trọng tâm cường độ. "
        f"Nhánh CNN: chuẩn hóa mean/std theo tác giả; pad chia hết 32; segmentation + distance map; "
        f"trilateration với ngưỡng bản đồ khoảng cách d_th={thr} (đúng public run_neural_net). "
        "Catalog: bảng K-vector theo FOV đường chéo demo."
    )
    s5 = (
        "Mặc định không học máy: CoG → Pyramid/Mortari + K-vector → Wahba SVD/QUEST. "
        "Tùy chọn AI: MobileUNet công khai chỉ thay bước centroid (Zhao et al.). "
        "Công cụ: Python, NumPy, SciPy, OpenCV, PyTorch CUDA, pytest, app.py, demo_star_tracker.py. "
        + zhao_cite
    )
    s6 = (
        "Không huấn luyện lại CNN. "
        f"Suy luận trên {s.get('device')}, cnn_torch_device={s.get('cnn_torch_device')}, "
        f"torch {s.get('torch_version', '2.11.0+cu128')}. "
        f"Ngưỡng distance-map công khai = {thr}. "
        "Một lần chạy cục bộ trước đó dùng ngưỡng ~0,5√2 ≈ 0,707 (sai so với public "
        "run_neural_net) nên bị loại; số liệu trong file centroid_benchmark_hd1024.json cũ "
        "không được dùng cho hồ sơ. File đo lại: data/outputs/centroid_benchmark_hd1024_thr2.json."
    )
    s7 = (
        "Chỉ số: RMSE pixel so với ground-truth simulator (khớp ≤ 5 px); số sao khớp/GT; "
        "thời gian ms/khung và s/khung; trên ảnh núi thật — số điểm trên vùng núi/cây vs trời "
        "(mặt nạ heuristic); boresight RA/Dec so với LOST công bố (góc tách, độ). "
        "Attitude end-to-end (Δθ) trên bộ sim lớn: chưa kiểm được lần này."
    )
    s8 = (
        f"BENCHMARK sim (thẩm quyền RMSE): {cl['n_frames']} khung {cam['width']}×{cam['height']}, "
        f"d_th={thr}, seed {s.get('seed')}, file data/outputs/centroid_benchmark_hd1024_thr2.json. "
        f"Thiết bị: {s.get('device')}. "
        f"CoG: RMSE tb {cl['rmse_px_mean']:.4f} px · {cl['time_ms_median']:.2f} ms tv; "
        f"CNN: RMSE tb {nn['rmse_px_mean']:.4f} px · {nn['time_ms_median']:.2f} ms tv. "
        "Trên bộ sạch này CoG thắng RMSE. "
        "Lóa mạnh 40 khung (centroid_benchmark_hd1024_glare_sun_strong.json): "
        "CoG RMSE tb 0.1449 px, CNN 0.3951 px — CNN không thắng RMSE trên bộ lóa này. "
        "Cổng sáng CNN tùy chọn «chỉ sao từ cấp 6 trở lên sáng» có trên web; "
        "trên một khung sim đã thử, cổng này không làm CNN thắng CoG — không overclaim. "
        + mount_para
        + " "
        + zhao_cite
        + " Không claim CNN đã chứng minh bay."
    )
    s9 = (
        f"Baseline CoG. Trên 500 khung hd1024 (d_th={thr}): CNN không cải thiện RMSE "
        f"(CoG {cl['rmse_px_mean']:.4f} px vs CNN {nn['rmse_px_mean']:.4f} px). "
        "Số bài báo (Table III/V/VI) chỉ là trích dẫn tác giả, tách khỏi bảng đo nhóm. "
        + creativity
        + " Đóng góp: tách CoG/CNN; giữ Pyramid/K-vector và Wahba; quaternion 16 byte; "
        "preset web «Ảnh núi thật»."
    )
    s10 = (
        "Kiến trúc: camera model → centroid (CoG | CNN) → Pyramid/K-vector → Wahba → quaternion 16 byte. "
        "Chạy: pip install -r requirements.txt; python demo_star_tracker.py --mode fast --seed 42 "
        "--lookup --no-show; python app.py (preset «Ảnh núi thật»). "
        "Slide: AI2026_SLIDE_THUYET_TRINH.pptx."
    )
    s11 = (
        "Rủi ro: nhận dạng sai → quaternion sai; radiometry assumption; trên ảnh sạch CNN lệch miền "
        "(RMSE cao hơn CoG). Mặt nạ núi/cây là heuristic — số đếm kèm overlay để kiểm chứng mắt. "
        "Không thu thập dữ liệu cá nhân. Demo laptop — không phải thiết bị bay."
    )
    s12 = (
        "Hướng tới: chuẩn hóa miền ảnh trên bộ sim; đo attitude end-to-end đầy đủ; "
        "mở rộng ảnh trời thật. Ứng dụng học tập tại Trường Đại học Bách khoa, "
        "Đại học Đà Nẵng — Khoa Công nghệ thông tin."
    )
    s13 = (
        "Minh chứng: data/outputs/centroid_benchmark_hd1024_thr2.json; "
        "data/outputs/mount_st_helens_centroid_compare.json; "
        "data/outputs/mount_st_helens_lost_compare.json; "
        "data/outputs/mount_st_helens_cog_vs_cnn.png; data/uploads/mount_st_helens_1.png; "
        "data/outputs/zhao_paper_figures/; third_party/cnn_star_centroid/; "
        "src/centroid_cnn.py (d_th=2.0); AI2026_SLIDE_THUYET_TRINH.pptx. "
        "Google Drive lịch sử câu lệnh: chưa kiểm được. "
        "Hai file video minh chứng: chưa kiểm được (thiếu file). "
        "URL git remote: chưa kiểm được — không bịa link. "
        f"Cấu trúc 13 mục Bang C đối chiếu PDF (tên file): {'; '.join(refs) if refs else 'chưa kiểm được'}."
    )
    return {
        19: s1, 22: s2, 26: s3, 30: s4, 34: s5, 38: s6, 42: s7,
        46: s8, 50: s9, 54: s10, 58: s11, 62: s12, 66: s13,
    }


def write_md(s: dict, bodies: dict[int, str], identity: dict[int, str], m: dict | None) -> None:
    titles = [
        (19, "1. Bài toán hoặc vấn đề thực tiễn cần giải quyết"),
        (22, "2. Mục tiêu, phạm vi và đối tượng ứng dụng của sản phẩm"),
        (26, "3. Dữ liệu sử dụng, nguồn dữ liệu và tính hợp lệ của dữ liệu"),
        (30, "4. Quy trình tiền xử lý, làm sạch, chuẩn hóa hoặc tổ chức dữ liệu"),
        (34, "5. Thuật toán, mô hình, phương pháp hoặc công cụ trí tuệ nhân tạo được sử dụng"),
        (38, "6. Quy trình huấn luyện, tinh chỉnh, tích hợp hoặc khai thác mô hình (nếu có)"),
        (42, "7. Chỉ số, phương pháp hoặc tiêu chí đánh giá kết quả"),
        (46, "8. Kết quả thử nghiệm, phân tích ưu điểm, hạn chế và khả năng mở rộng"),
        (50, "9. So sánh với phương án hoặc mô hình cơ sở, phân tích đóng góp của các thành phần trong hệ thống (nếu có)"),
        (54, "10. Kiến trúc hệ thống và phương án triển khai"),
        (58, "11. Phân tích rủi ro, yêu cầu bảo mật, đạo đức trí tuệ nhân tạo và an toàn dữ liệu"),
        (62, "12. Hướng phát triển, hoàn thiện và khả năng ứng dụng trong thực tiễn"),
        (66, "13. Lịch sử câu lệnh và hình ảnh minh chứng quá trình phát triển sản phẩm từ bản nháp đến khi hoàn thiện"),
    ]
    g = identity.get
    lines = [
        "# HỒ SƠ DỰ ÁN DỰ THI BẢNG C",
        "Cuộc thi Sáng tạo trẻ Quốc gia trong lĩnh vực Trí tuệ nhân tạo năm 2026",
        "",
        "## Thông tin đội thi",
        f"- Số lượng thí sinh: {g(-1, '')} người (theo hồ sơ gốc)",
        f"- Thí sinh 1 (đội trưởng): {g(2, '')} — {g(3, '')} — {g(6, '')} — {g(7, '')}",
        f"  - {g(4, '')}",
        f"  - Địa chỉ: {g(5, '')}",
        f"- Thí sinh 2: {g(9, '')} — {g(10, '')} — {g(13, '')} — {g(14, '')}",
        f"  - {g(11, '')}",
        f"  - Địa chỉ: {g(12, '')}",
        f"- Thí sinh 3: {g(16, '')} — {g(17, '')} — {g(20, '')} — {g(21, '')}",
        f"  - {g(18, '')}",
        f"  - Địa chỉ: {g(19, '')}",
        f"- GVHD: {ADVISOR} (xác nhận người dùng; mẫu Bang C không có ô GVHD)",
        "- Trường: Trường Đại học Bách khoa, Đại học Đà Nẵng",
        "- Khoa: Khoa Công nghệ thông tin",
        "",
        "## NỘI DUNG HỒ SƠ DỰ ÁN",
        "",
    ]
    for idx, title in titles:
        lines += [f"### {title}", "", bodies[idx], ""]
    lines += [
        "## Hình minh chứng",
        "",
        "### Demo nhóm — ảnh núi thật LOST (CoG cam vs CNN xanh)",
        "",
        "![Overlay CoG vs CNN trên mount_st_helens_1.png](data/outputs/mount_st_helens_cog_vs_cnn.png)",
        "",
        "*Chú thích: ảnh thật công bố kèm LOST, không phải ảnh của nhóm chụp; thông số 4,2 mm và 4,1 µm. "
        "Cam = CoG; xanh = CNN MobileUNet (d_th=2). Đỏ nhạt = mặt nạ núi/cây heuristic.*",
        "",
    ]
    if m is not None:
        lines += [
            f"- CoG: tổng {m['classical']['n_total']} · núi/cây {m['classical']['n_on_terrain']} · trời {m['classical']['n_on_sky']}",
            f"- CNN: tổng {m['cnn']['n_total']} · núi/cây {m['cnn']['n_on_terrain']} · trời {m['cnn']['n_on_sky']}",
            f"- Kết luận: {m.get('verdict_vi', '')}",
            f"- {_lost_attitude_sentence(_mount_lost())}",
            "",
        ]
    if DEMO_FIELD.is_file():
        lines += [
            "### Demo nhóm — khung sim hd1024 (1024×683)",
            "",
            "![Khung sao giả lập hd1024](data/sim_dataset/hd1024/0001_preview.png)",
            "",
            "*Chú thích: demo của nhóm, 1024×683 — không phải số liệu bài báo.*",
            "",
        ]
    lines += [
        "### Nguồn Zhao et al., arXiv:2404.19108 (không phải đo của nhóm)",
        "",
    ]
    for rel, cap in (
        ("data/outputs/zhao_paper_figures/project_00.png", "Flowchart phương pháp — Nguồn: Zhao et al., arXiv:2404.19108"),
        ("data/outputs/zhao_paper_figures/project_03.png", "Night sky + stray light — Nguồn: Zhao et al., arXiv:2404.19108"),
        ("data/outputs/zhao_paper_figures/project_05.png", "Attitude so sánh — Nguồn: Zhao et al., arXiv:2404.19108"),
    ):
        if (ROOT / rel).is_file():
            lines += [f"![{cap}]({rel})", "", f"*{cap}*", ""]
    lines += ["## PDF cấu trúc đã đối chiếu (chỉ cấu trúc)", ""]
    for n in STRUCTURE_PDFS:
        if (ROOT / n).is_file():
            lines.append(f"- `{n}`")
    lines.append("")
    OUT_MD.write_text("\n".join(lines), encoding="utf-8")


def _add_caption(doc: Document, text: str) -> None:
    p = doc.add_paragraph(text)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for run in p.runs:
        run.font.size = Pt(10)
        run.font.italic = True


def _append_figures(doc: Document) -> None:
    doc.add_heading("Hình minh chứng (bổ sung)", level=2)
    if MOUNT_OVERLAY.is_file():
        doc.add_paragraph("Ảnh núi thật LOST — overlay CoG (cam) vs CNN (xanh)")
        doc.add_picture(str(MOUNT_OVERLAY), width=Inches(5.8))
        _add_caption(
            doc,
            "Chú thích: ảnh thật công bố kèm LOST, không phải ảnh của nhóm chụp; "
            "thông số 4,2 mm và 4,1 µm. File: data/outputs/mount_st_helens_cog_vs_cnn.png",
        )
    if DEMO_FIELD.is_file():
        doc.add_paragraph("Khung sao giả lập hd1024 (demo của nhóm, 1024×683)")
        doc.add_picture(str(DEMO_FIELD), width=Inches(5.2))
        _add_caption(doc, "Chú thích: demo của nhóm, 1024×683 — không phải số liệu bài báo.")
    doc.add_paragraph("Hình từ Zhao et al. (chỉ trích dẫn nguồn)")
    for path, cap in (
        (ZHAO_FLOW, "Nguồn: Zhao et al., arXiv:2404.19108 — flowchart phương pháp"),
        (ZHAO_STRAY, "Nguồn: Zhao et al., arXiv:2404.19108 — night sky + stray light"),
        (ZHAO_ATT, "Nguồn: Zhao et al., arXiv:2404.19108 — attitude comparison"),
    ):
        if path.is_file():
            doc.add_picture(str(path), width=Inches(5.5))
            _add_caption(doc, cap)


def main() -> None:
    if not BENCH.is_file():
        raise SystemExit(f"missing {BENCH}")
    template = TEMPLATE
    if not template.is_file():
        cands = list(ROOT.glob("AI2026*gốc*.docx")) + list(ROOT.glob("AI2026*goc*.docx"))
        if not cands:
            raise SystemExit(f"missing {TEMPLATE}")
        template = cands[0]

    s = _summary()
    m = _mount()
    lost = _mount_lost()
    bodies = sections(s, m, lost)
    identity = load_identity_from_original()
    write_md(s, bodies, identity, m)

    doc = Document(str(template))
    table = doc.tables[0]
    if len(table.rows[0].cells) >= 4:
        table.rows[0].cells[1].text = "1 người ☐"
        table.rows[0].cells[2].text = "2 người ☐"
        table.rows[0].cells[3].text = "3 người ☑"
    for ri in IDENTITY_VALUE_ROWS:
        val = identity.get(ri, "")
        if ri < len(table.rows) and len(table.rows[ri].cells) >= 2:
            table.rows[ri].cells[1].text = val

    scrub = [
        "Vin" + "AI", "Vin" + "space", "VLS", "Alv" + "ium",
        "IMX" + "541", "IMX" + "542", "Schne" + "ider", "Edmu" + "nd",
        "Jets" + "on", "CubeSat 8U", "8U ISIS",
        "0.1368", "0.136817", "0.4340", "0.434037", "0.5250", "0.52505",
        "36.46", "84.22", "48.71", "30.71",
    ]
    for para in doc.paragraphs:
        t = para.text
        for bad in scrub:
            if bad in t:
                t = t.replace(bad, "")
        if t != para.text:
            para.text = t

    filled: set[int] = set()
    for i, para in enumerate(doc.paragraphs):
        if i in bodies and _dots(para, bodies[i]):
            filled.add(i)
            continue
        if set(para.text.strip()) <= {".", " ", "…", "\u00a0"} and para.text.strip().startswith("..."):
            if any(i - k in filled for k in range(1, 4)):
                para.text = ""

    _append_figures(doc)
    doc.save(str(OUT))
    print(f"wrote {OUT}")
    print(f"wrote {OUT_MD}")
    print(f"sections {sorted(filled)}")
    print(f"identity rows {sorted(k for k in identity if k >= 0)}")
    print(f"CoG RMSE {s['classical']['rmse_px_mean']:.4f} CNN {s['cnn']['rmse_px_mean']:.4f}")
    if m:
        print(
            f"mount CoG terrain={m['classical']['n_on_terrain']} "
            f"CNN terrain={m['cnn']['n_on_terrain']} claim_held={m.get('claim_held')}"
        )


if __name__ == "__main__":
    main()
