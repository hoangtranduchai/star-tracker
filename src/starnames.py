"""IAU proper-name overlay. ADCS core still uses catalog indices, not names.

Primary source: WGSN "IAU-Catalog of Star Names (always up to date).xlsx"
(the file the user downloaded into star_tracker/data).
Fallback: 2022-04-04 IAU-CSN.txt (HR-only dump).
"""

from __future__ import annotations

import json
import re
from functools import lru_cache
from pathlib import Path

ROOT = Path(__file__).resolve().parent.parent
IAU_TXT = ROOT / "data" / "catalogs" / "IAU-CSN.txt"
IAU_JSON = ROOT / "data" / "catalogs" / "iau_hr_names.json"
IAU_XLSX_NAME = "IAU-Catalog of Star Names (always up to date).xlsx"
HR_PAT = re.compile(r"HR\s*(\d+)\b", re.I)
JSON_FORMAT = "iau_csn_v2"


def find_iau_xlsx() -> Path | None:
    exact = ROOT / "data" / IAU_XLSX_NAME
    if exact.is_file():
        return exact
    hits = sorted(ROOT.joinpath("data").glob("**/*Star Names*.xlsx"))
    return hits[0] if hits else None


def _hip_to_hr_from_txt(path: Path) -> dict[int, int]:
    out: dict[int, int] = {}
    if not path.is_file():
        return out
    for line in path.read_text(encoding="utf-8", errors="replace").splitlines():
        s = line.strip()
        if not s or s[0] in "#$":
            continue
        m = re.search(r"\sHR\s+(\d+)\b", " " + s)
        if not m:
            continue
        hr = int(m.group(1))
        rest = s[m.end() :].split()
        for i, tok in enumerate(rest):
            if tok in {"V", "G"} and i + 1 < len(rest) and rest[i + 1].isdigit():
                out[int(rest[i + 1])] = hr
                break
    return out


def parse_iau_txt(path: Path) -> dict[int, dict]:
    out: dict[int, dict] = {}
    text = path.read_text(encoding="utf-8", errors="replace")
    for line in text.splitlines():
        s = line.strip()
        if not s or s.startswith("#") or s.startswith("$"):
            continue
        m = re.search(r"\sHR\s+(\d+)\b", " " + s)
        if not m:
            continue
        hr = int(m.group(1))
        left = s[: m.start()].strip()
        words = left.split()
        name = words[0] if words else f"HR{hr}"
        for k in range(1, len(words) + 1):
            if 2 * k == len(words) and words[:k] == words[k:]:
                name = " ".join(words[:k])
                break
        rest = s[m.end() :].strip().split()
        designation = rest[0] if rest else ""
        if designation in {"_", "-"}:
            designation = ""
        out[hr] = {"name": name, "designation": designation, "hr": hr}
    return out


def _as_int(value) -> int | None:
    if value is None:
        return None
    if isinstance(value, bool):
        return None
    if isinstance(value, int):
        return value
    text = str(value).strip()
    if not text or text in {"_", "-"}:
        return None
    try:
        return int(float(text.split()[0]))
    except (TypeError, ValueError):
        return None


def parse_iau_xlsx(path: Path | None = None) -> list[dict]:
    """Rows from the WGSN Excel CSN. HR filled from Designation, else HIP/name vs 2022 txt."""
    import openpyxl

    path = Path(path) if path is not None else find_iau_xlsx()
    if path is None or not path.is_file():
        raise FileNotFoundError("IAU CSN xlsx not found under star_tracker/data")

    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    sheet = wb[wb.sheetnames[0]]
    raw_rows = list(sheet.iter_rows(values_only=True))
    if len(raw_rows) < 3:
        return []
    header_idx = next(
        i for i, row in enumerate(raw_rows) if row and row[0] == "proper names"
    )
    header = [str(c).strip() if c is not None else "" for c in raw_rows[header_idx]]
    col = {name: i for i, name in enumerate(header)}
    old_hr = parse_iau_txt(IAU_TXT) if IAU_TXT.is_file() else {}
    name_to_hr = {rec["name"].lower(): hr for hr, rec in old_hr.items()}
    hip_to_hr = _hip_to_hr_from_txt(IAU_TXT)

    rows: list[dict] = []
    for row in raw_rows[header_idx + 1 :]:
        if not row or row[0] is None or not str(row[0]).strip():
            continue
        name = str(row[col["proper names"]]).strip()
        des = ""
        if "Designation" in col and col["Designation"] < len(row) and row[col["Designation"]] is not None:
            des = str(row[col["Designation"]]).strip()
        hip = _as_int(row[col["HIP"]]) if "HIP" in col and col["HIP"] < len(row) else None
        bayer = ""
        if "Bayer ID" in col and col["Bayer ID"] < len(row) and row[col["Bayer ID"]] is not None:
            bayer = str(row[col["Bayer ID"]]).strip()
        hr = None
        m = HR_PAT.search(des)
        if m:
            hr = int(m.group(1))
        elif hip is not None and hip in hip_to_hr:
            hr = hip_to_hr[hip]
        elif name.lower() in name_to_hr:
            hr = name_to_hr[name.lower()]
        rec = {
            "name": name,
            "designation": des or bayer,
            "hr": hr,
            "hip": hip,
        }
        rows.append(rec)
    return rows


def _records_to_bundle(rows: list[dict], source: str) -> dict:
    by_hr: dict[str, dict] = {}
    by_hip: dict[str, dict] = {}
    for rec in rows:
        if rec.get("hr") is not None:
            by_hr[str(int(rec["hr"]))] = rec
        if rec.get("hip") is not None:
            by_hip[str(int(rec["hip"]))] = rec
    return {
        "format": JSON_FORMAT,
        "source": source,
        "n_names": len(rows),
        "by_hr": by_hr,
        "by_hip": by_hip,
    }


def _legacy_bundle(raw: dict) -> dict:
    by_hr = {str(int(k)): v for k, v in raw.items()}
    return {
        "format": JSON_FORMAT,
        "source": "legacy-json",
        "n_names": len(by_hr),
        "by_hr": by_hr,
        "by_hip": {},
    }


def _write_bundle(txt: Path | None = None) -> dict:
    xlsx = find_iau_xlsx()
    if xlsx is not None:
        bundle = _records_to_bundle(parse_iau_xlsx(xlsx), xlsx.name)
    else:
        parsed = parse_iau_txt(txt or IAU_TXT)
        rows = [
            {"name": v["name"], "designation": v.get("designation") or "", "hr": k, "hip": None}
            for k, v in parsed.items()
        ]
        bundle = _records_to_bundle(rows, (txt or IAU_TXT).name)
    IAU_JSON.parent.mkdir(parents=True, exist_ok=True)
    IAU_JSON.write_text(json.dumps(bundle, indent=2, ensure_ascii=False), encoding="utf-8")
    return bundle


def rebuild_cache(txt: Path | None = None) -> Path:
    _write_bundle(txt)
    _bundle.cache_clear()
    return IAU_JSON


def _xlsx_newer_than_json() -> bool:
    xlsx = find_iau_xlsx()
    if xlsx is None or not IAU_JSON.exists():
        return xlsx is not None
    return xlsx.stat().st_mtime > IAU_JSON.stat().st_mtime


@lru_cache(maxsize=1)
def _bundle() -> dict:
    if _xlsx_newer_than_json() or not IAU_JSON.exists():
        return _write_bundle()
    raw = json.loads(IAU_JSON.read_text(encoding="utf-8"))
    if isinstance(raw, dict) and raw.get("format") == JSON_FORMAT:
        return raw
    if isinstance(raw, dict) and "by_hr" not in raw:
        return _legacy_bundle(raw)
    return raw


def _table() -> dict[int, dict]:
    return {int(k): v for k, v in _bundle()["by_hr"].items()}


def lookup_hr(hr: int) -> dict | None:
    return _table().get(int(hr))


def lookup_hip(hip: int) -> dict | None:
    rec = _bundle()["by_hip"].get(str(int(hip)))
    return rec


def n_named_hr() -> int:
    return len(_bundle()["by_hr"])


def n_iau_names() -> int:
    return int(_bundle()["n_names"])


def label_hr(hr: int, compact: bool = True) -> str:
    rec = lookup_hr(hr)
    if rec and rec.get("name"):
        if compact:
            return f"{rec['name']} HR{int(hr)}"
        des = rec.get("designation") or ""
        extra = f" ({des})" if des else ""
        return f"{rec['name']}{extra}  HR {int(hr)}"
    return f"HR {int(hr)}"


def id_scheme(catalog_source: str) -> str:
    src = str(catalog_source).upper()
    if "BSC5" in src or "BRIGHT STAR" in src:
        return "HR"
    if "HIP" in src:
        return "HIP"
    return "ID"


def describe_index(catalog, idx: int) -> dict:
    """Human-readable row for one catalog index. Not used by Pyramid/Wahba."""
    import numpy as np

    from .catalog import vectors_to_ra_dec

    idx = int(idx)
    star_id = int(catalog.ids[idx])
    scheme = id_scheme(catalog.source)
    if scheme == "HR":
        rec = lookup_hr(star_id)
        label = label_hr(star_id, compact=False) if rec and rec.get("name") else f"HR {star_id}"
    elif scheme == "HIP":
        rec = lookup_hip(star_id)
        if rec and rec.get("name"):
            label = f"{rec['name']}  HIP {star_id}"
        else:
            rec = None
            label = f"HIP {star_id}"
    else:
        rec = None
        label = f"{scheme} {star_id}"
    ra, dec = vectors_to_ra_dec(catalog.vectors[idx : idx + 1])
    return {
        "catalog_idx": idx,
        "id": star_id,
        "id_scheme": scheme,
        "iau_name": rec.get("name") if rec else None,
        "designation": (rec.get("designation") or None) if rec else None,
        "label": label,
        "mag": float(catalog.mags[idx]),
        "ra_deg": float(np.degrees(ra[0])),
        "dec_deg": float(np.degrees(dec[0])),
    }


def compact_tag(catalog, idx: int, lookup: bool) -> str:
    star_id = int(catalog.ids[int(idx)])
    scheme = id_scheme(catalog.source)
    if lookup and scheme == "HR":
        return label_hr(star_id, compact=True)
    if lookup and scheme == "HIP":
        rec = lookup_hip(star_id)
        if rec and rec.get("name"):
            return f"{rec['name']} HIP{star_id}"
    return f"{scheme} {star_id}"
